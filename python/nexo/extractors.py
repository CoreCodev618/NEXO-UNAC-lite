"""Extraccion multi-formato para el pipeline RAG de NEXO-UNAC (fase IA-3).

Formatos soportados:
- .docx   (Word)                 -> texto de parrafos y tablas
- .pptx   (PowerPoint)           -> texto por diapositiva + notas del orador
- .pdf                           -> texto por pagina (si esta escaneado, avisa)
- .jpg/.jpeg/.png/.webp          -> OCR con Gemini Vision (pizarras / examenes)
- archivos/carpetas de codigo    -> se ignora system/, vendor/, writable/, node_modules/
- .zip                           -> descompresion controlada y se procesa adentro

Seguridad (zip): rechaza rutas con "..", rutas absolutas y symlinks; limita el
tamano total descomprimido (200 MB), el numero de entradas (2000) y el tamano
por archivo de codigo (512 KB). Nunca se ejecuta nada de lo que venga adentro.
"""
from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from tempfile import TemporaryDirectory

from . import ocr

MAX_ZIP_BYTES = 200 * 1024 * 1024
MAX_ZIP_FILES = 2000
MAX_CODE_FILE_BYTES = 512 * 1024
MAX_FILES_PER_DIR = 400

SKIP_DIRS = {
    "system", "vendor", "node_modules", "writable", ".git",
    "__pycache__", ".venv", "ThirdParty", "build", "dist",
}

CODE_EXT = {
    ".php", ".js", ".ts", ".jsx", ".tsx", ".html", ".css", ".scss",
    ".sql", ".py", ".java", ".c", ".cpp", ".h", ".cs", ".rb", ".go",
    ".rs", ".kt", ".swift", ".sh", ".md", ".txt",
}
IMG_EXT = {".jpg", ".jpeg", ".png", ".webp"}
DOC_EXT = {".docx", ".pptx", ".pdf"}


@dataclass
class Unit:
    """Trozo de texto extraido con su ubicacion dentro de la fuente."""
    location: str
    text: str


@dataclass
class Extracted:
    """Una fuente procesada (archivo) con sus unidades de texto."""
    source: str
    kind: str
    units: list[Unit] = field(default_factory=list)
    # Identidad canonica (ruta relativa al proyecto) para detectar colisiones.
    source_key: str = ""
    # documents.id de MySQL cuando el ingest recibe --doc-id/--doc-id-map.
    doc_id: int | None = None

    @property
    def chars(self) -> int:
        return sum(len(u.text) for u in self.units)


@dataclass
class Extraction:
    """Resultado de extraer un archivo o carpeta."""
    items: list[Extracted] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def extract_path(target: Path) -> Extraction:
    """Punto de entrada: extrae un archivo, carpeta o .zip."""
    res = Extraction()
    if target.is_dir():
        _extract_dir(target, res)
    else:
        _extract_file(target, res, source=target.name)
    return res


# ---------------------------------------------------------------------------
# Internos


def _ignored(path: Path) -> bool:
    return any(part in SKIP_DIRS for part in path.parts)


def _extract_dir(root: Path, res: Extraction) -> None:
    files = [p for p in sorted(root.rglob("*")) if p.is_file() and not _ignored(p)]
    if len(files) > MAX_FILES_PER_DIR:
        res.warnings.append(
            f"{root.name}: corte a {MAX_FILES_PER_DIR} archivos (habia {len(files)})"
        )
        files = files[:MAX_FILES_PER_DIR]

    skipped: dict[str, int] = {}
    for path in files:
        ext = path.suffix.lower()
        rel = path.relative_to(root)
        if ext in DOC_EXT or ext in IMG_EXT or ext in CODE_EXT or ext == ".zip":
            _extract_file(path, res, source=f"{root.name}/{rel}")
        else:
            key = ext or "(sin ext)"
            skipped[key] = skipped.get(key, 0) + 1
    if skipped:
        detail = ", ".join(f"{n} {e}" for e, n in sorted(skipped.items(), key=lambda x: -x[1])[:6])
        res.warnings.append(f"{root.name}: omitidos {sum(skipped.values())} archivos no soportados ({detail})")


def _extract_file(path: Path, res: Extraction, source: str) -> None:
    ext = path.suffix.lower()
    if ext == ".zip":
        _extract_zip(path, res, source)
        return

    try:
        if ext == ".docx":
            kind, units = "docx", _docx_units(path)
        elif ext == ".pptx":
            kind, units = "pptx", _pptx_units(path)
        elif ext == ".pdf":
            kind, units = "pdf", _pdf_units(path)
        elif ext in IMG_EXT:
            kind, units = "imagen", [Unit("imagen", ocr.ocr_image(path))]
        elif ext in CODE_EXT:
            kind, units = "codigo", _code_units(path, res, source)
        else:
            return
    except Exception as exc:  # noqa: BLE001 - queremos seguir con el resto
        res.warnings.append(f"{source}: ERROR {type(exc).__name__}: {exc}")
        return

    if units and any(u.text.strip() for u in units):
        res.items.append(Extracted(source=source, kind=kind, units=units))
    else:
        res.warnings.append(f"{source}: sin texto extraido")


def _extract_zip(path: Path, res: Extraction, source: str) -> None:
    with zipfile.ZipFile(path) as zf:
        infos = [i for i in zf.infolist() if not i.is_dir()]
        total = sum(i.file_size for i in infos)
        if total > MAX_ZIP_BYTES:
            res.warnings.append(
                f"{source}: zip demasiado grande descomprimido ({total / 1e6:.1f} MB > 200 MB), omitido"
            )
            return
        if len(infos) > MAX_ZIP_FILES:
            res.warnings.append(f"{source}: zip con demasiadas entradas ({len(infos)}), omitido")
            return

        with TemporaryDirectory(prefix="nexo_zip_") as tmp:
            base = Path(tmp)
            for info in infos:
                rel = Path(info.filename.replace("\\", "/"))
                if rel.is_absolute() or ".." in rel.parts:
                    res.warnings.append(f"{source}: entrada con ruta sospechosa ignorada ({info.filename})")
                    continue
                if (info.external_attr >> 16) & 0o170000 == 0o120000:  # symlink
                    res.warnings.append(f"{source}: symlink ignorado ({info.filename})")
                    continue
                zf.extract(info, base)

            inner = [p for p in sorted(base.rglob("*")) if p.is_file() and not _ignored(p)]
            for i, inner_path in enumerate(inner[:MAX_FILES_PER_DIR], start=1):
                rel = inner_path.relative_to(base)
                _extract_file(inner_path, res, source=f"{source} :: {rel}")
            if len(inner) > MAX_FILES_PER_DIR:
                res.warnings.append(
                    f"{source}: corte a {MAX_FILES_PER_DIR} archivos internos (habia {len(inner)})"
                )


def _docx_units(path: Path) -> list[Unit]:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []

    try:
        blocks = list(doc.iter_inner_content())
    except AttributeError:  # python-docx viejo
        blocks = [*doc.paragraphs, *doc.tables]

    for block in blocks:
        if hasattr(block, "rows"):  # tabla
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        else:  # parrafo
            text = block.text.strip()
            if text:
                parts.append(text)

    text = "\n\n".join(parts).strip()
    return [Unit("documento", text)] if text else []


def _pptx_units(path: Path) -> list[Unit]:
    from pptx import Presentation

    prs = Presentation(str(path))
    units: list[Unit] = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts = _shape_texts(slide.shapes)
        if texts:
            units.append(Unit(f"slide {idx}", "\n".join(texts)))
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    units.append(Unit(f"slide {idx} (notas)", notes))
        except Exception:  # noqa: BLE001 - notas opcionales
            pass
    return units


def _shape_texts(shapes) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out: list[str] = []
    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                out.extend(_shape_texts(shape.shapes))
                continue
        except Exception:  # noqa: BLE001
            pass
        try:
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        out.append(" | ".join(cells))
                continue
        except Exception:  # noqa: BLE001
            pass
        try:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    out.append(text)
        except Exception:  # noqa: BLE001
            pass
    return out


def _pdf_units(path: Path) -> list[Unit]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    units: list[Unit] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            text = "[pagina sin texto extraible: probablemente escaneada; requiere OCR de PDF]"
        units.append(Unit(f"pagina {idx}", text))
    return units


def _code_units(path: Path, res: Extraction, source: str) -> list[Unit]:
    data = path.read_bytes()
    if len(data) > MAX_CODE_FILE_BYTES:
        res.warnings.append(f"{source}: archivo grande ({len(data) // 1024} KB > 512 KB), omitido")
        return []
    if b"\x00" in data[:2048]:
        res.warnings.append(f"{source}: parece binario, omitido")
        return []
    text = data.decode("utf-8", errors="replace").strip()
    return [Unit("archivo", text)] if text else []
